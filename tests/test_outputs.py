import json
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
TABLES = ROOT / "outputs" / "tables"
WORKBOOK = ROOT / "outputs" / "IFRS_9_Mortgage_ECL_Model.xlsx"


def table(name):
    return pd.read_csv(TABLES / f"{name}.csv")


def test_all_published_controls_pass():
    controls = table("controls")
    assert len(controls) >= 30
    assert controls["status"].eq("PASS").all()


def test_all_three_stages_and_methods_exist():
    account = table("account_level_ecl")
    traces = table("worked_trace_monthly")
    assert set(account["stage"]) == {1, 2, 3}
    assert traces.loc[traces["stage"].eq(3), "calculation_method"].str.contains("cash shortfall").all()


def test_scenario_weighting_and_portfolio_reconciliation():
    scenario = table("scenario_summary")
    account = table("account_level_ecl")
    stages = table("stage_summary")
    assert np.isclose(scenario["scenario_weight"].sum(), 1.0)
    assert np.isclose(scenario["weighted_contribution"].sum(), account["loan_ecl"].sum(), atol=0.01)
    assert np.isclose(account["loan_ecl"].sum(), stages["ecl"].sum(), atol=0.01)


def test_worked_projection_has_required_audit_fields():
    required = {
        "loan_id", "stage", "scenario", "future_month", "conditional_pd",
        "survival_probability", "marginal_pd", "cumulative_pd", "ead",
        "projected_property_value", "projected_ltv", "lgd", "discount_factor", "period_ecl",
    }
    assert required <= set(table("worked_trace_monthly").columns)


def test_stage_horizons_in_worked_traces():
    traces = table("worked_trace_monthly")
    horizons = traces.groupby("stage")["future_month"].max()
    assert horizons.loc[1] <= 12
    assert horizons.loc[2] > 12


def test_monthly_ecl_equation_and_stage3_cash_shortfall():
    traces = table("worked_trace_monthly")
    performing = traces[traces["stage"].isin([1, 2])]
    difference = performing["period_ecl"] - (
        performing["marginal_pd"] * performing["lgd"]
        * performing["ead"] * performing["discount_factor"]
    )
    stage3 = traces[traces["stage"].eq(3)]
    stage3_difference = stage3["period_ecl"] - (
        stage3["ead"] - stage3["discounted_expected_recovery"]
    ).clip(lower=0)
    assert difference.abs().max() < 1e-8
    assert stage3_difference.abs().max() < 1e-8


def test_worked_trace_scenario_weighting_reconciles():
    traces = table("worked_trace_monthly")
    summary = table("worked_trace_summary").set_index("loan_id")
    scenario = traces.groupby(["loan_id", "scenario", "scenario_weight"], as_index=False)["period_ecl"].sum()
    scenario["weighted"] = scenario["period_ecl"] * scenario["scenario_weight"]
    calculated = scenario.groupby("loan_id")["weighted"].sum()
    assert np.allclose(calculated.sort_index(), summary.loc[calculated.index, "loan_ecl"].sort_index(), atol=1e-8)


def test_stage3_pd_is_not_reported_as_performing_pd():
    account = table("account_level_ecl")
    assert account.loc[account["stage"].eq(3), "current_pd_12m"].isna().all()


def test_scenario_ordering_and_public_aliases():
    scenario = table("scenario_summary").set_index("scenario")["scenario_ecl"]
    assert scenario["Upside"] <= scenario["Base"] <= scenario["Downside"]
    aliases = table("account_level_ecl")["loan_id"]
    assert aliases.str.fullmatch(r"RM\d{6}").all()


def test_notebooks_are_executed_without_errors():
    for path in sorted((ROOT / "notebooks").glob("*.ipynb")):
        notebook = json.loads(path.read_text(encoding="utf-8"))
        code = [cell for cell in notebook["cells"] if cell["cell_type"] == "code"]
        assert code and all(cell.get("execution_count") is not None for cell in code)
        assert not [output for cell in code for output in cell.get("outputs", []) if output.get("output_type") == "error"]


def test_workbook_matches_current_reporting_structure_and_numeric_types():
    workbook = load_workbook(WORKBOOK, data_only=False, read_only=False)
    required = {
        "Executive Dashboard", "PD Macro Model", "PD Term Structure", "LGD Term Structure",
        "EAD Term Structure", "Discount Factors", "Sources & Assumptions",
    }
    assert required <= set(workbook.sheetnames)
    assert "PD Macro" not in workbook.sheetnames
    for sheet, cells in {"Stage Summary": ["A2", "B2", "C2", "D2"], "Scenario Summary": ["B2", "C2", "D2"]}.items():
        assert all(isinstance(workbook[sheet][cell].value, (int, float)) for cell in cells)
    with zipfile.ZipFile(WORKBOOK) as package:
        xml = "".join(
            package.read(name).decode("utf-8", "ignore")
            for name in package.namelist() if name.endswith(".xml")
        )
    assert "ChatGPT" not in xml
    assert "OpenAI" not in xml
