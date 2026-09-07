from __future__ import annotations

from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SHEETS = [
    ("stage_summary.csv", "Stage Summary"),
    ("scenario_summary.csv", "Scenario Summary"),
    ("monthly_ecl_term_structure.csv", "Monthly ECL"),
    ("transition_matrix.csv", "Transition Matrix"),
    ("pd_distribution.csv", "PD Distribution"),
    ("lgd_distribution.csv", "LGD Distribution"),
    ("ead_profile.csv", "EAD Profile"),
    ("ltv_analysis.csv", "LTV Analysis"),
    ("lgd_by_ltv.csv", "LGD by LTV"),
    ("dpd_analysis.csv", "DPD Analysis"),
    ("geography_analysis.csv", "Geography"),
    ("lifetime_pd_distribution.csv", "Lifetime PD"),
    ("vintage_analysis.csv", "Vintage"),
    ("risk_grade_analysis.csv", "Risk Grade"),
    ("account_level_ecl.csv", "Account ECL"),
    ("macro_scenarios.csv", "Macro Scenarios"),
    ("pd_model_metrics.csv", "PD Metrics"),
    ("pd_woe_iv.csv", "PD WOE IV"),
    ("pd_calibration.csv", "PD Calibration"),
    ("lgd_model_metrics.csv", "LGD Metrics"),
    ("ead_backtest_metrics.csv", "EAD Backtest"),
    ("ead_macro_sensitivity.csv", "EAD Macro Test"),
    ("scenario_weight_analysis.csv", "Scenario Weights"),
    ("pd_macro_satellite_metrics.csv", "PD Macro Model"),
    ("pd_term_structure.csv", "PD Term Structure"),
    ("lgd_term_structure.csv", "LGD Term Structure"),
    ("ead_term_structure.csv", "EAD Term Structure"),
    ("discount_factor_term_structure.csv", "Discount Factors"),
    ("recovery_summary.csv", "Recovery Summary"),
    ("controls.csv", "Controls"),
    ("worked_trace_summary.csv", "Trace Summary"),
    ("worked_trace_monthly.csv", "Trace Monthly"),
]

NAVY = "17365D"
BLUE = "1F4E78"
PALE_BLUE = "D9EAF7"
PALE_GREEN = "E2F0D9"
PALE_AMBER = "FFF2CC"
WHITE = "FFFFFF"
GREEN = "006100"


def _number_format(header: str) -> str | None:
    h = header.lower()
    fraction_fields = (
        "scenario_weight", "coverage_ratio", "exposure_share", "ecl_share", "transition_probability",
        "current_pd_12m", "origination_pd_12m", "current_lifetime_pd", "origination_lifetime_pd",
        "marginal_pd", "conditional_pd", "survival_probability", "cumulative_pd", "lgd",
        "average_lgd", "observed_lgd", "predicted_lgd", "roc_auc", "gini", "ks", "brier_score",
        "observed_rate", "predicted_rate", "default_rate", "mae", "rmse", "r_squared",
        "mape", "weighted_absolute_percentage_error", "macro_auc_improvement",
        "configured_weight", "historical_analogue_frequency",
        "selected_weight",
    )
    point_fields = ("unemployment_rate", "gdp_growth_yoy", "hpi_growth_yoy", "mortgage_rate", "current_interest_rate", "current_ltv", "projected_ltv", "original_ltv")
    if h in fraction_fields:
        return "0.00%"
    if h.endswith("_usd"):
        return '"$"#,##0'
    currency_tokens = ("exposure", "ecl", "ead", "upb", "property_value", "contribution", "actual_loss")
    currency_names = {"net_sale_proceeds", "mi_recoveries", "other_recoveries", "recovery_expenses", "discounted_recoveries", "expected_recovery", "discounted_expected_recovery"}
    if any(token in h for token in currency_tokens) or h in currency_names:
        return '"$"#,##0'
    if h == "period" or h.endswith("_date") or "projection_period" in h or "reporting_period" in h:
        return "yyyy-mm-dd"
    if any(field in h for field in point_fields):
        return "0.00"
    if h in ("discount_factor", "hpi_factor", "woe", "iv_component", "variable_iv", "calibration_ratio"):
        return "0.0000"
    return None


def _safe_value(value):
    if pd.isna(value):
        return None
    if hasattr(value, "to_pydatetime"):
        return value.to_pydatetime()
    if isinstance(value, (pd.Interval,)):
        return str(value)
    if hasattr(value, "item"):
        return value.item()
    return value


def _write_frame(ws, frame: pd.DataFrame) -> None:
    headers = list(frame.columns)
    ws.append(headers)
    for row in frame.itertuples(index=False, name=None):
        ws.append([_safe_value(value) for value in row])
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor=NAVY)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for col_index, header in enumerate(headers, start=1):
        fmt = _number_format(header)
        if fmt is None and pd.api.types.is_integer_dtype(frame[header].dtype):
            fmt = "#,##0"
        elif fmt is None and pd.api.types.is_numeric_dtype(frame[header].dtype):
            fmt = "#,##0.000000"
        max_len = len(str(header))
        for cell in ws.iter_cols(min_col=col_index, max_col=col_index, min_row=2, max_row=ws.max_row):
            for item in cell:
                if fmt:
                    item.number_format = fmt
                item.font = Font(name="Aptos", size=10, color="222222")
                if item.row <= 200:
                    max_len = max(max_len, len(str(item.value)) if item.value is not None else 0)
        width = min(max(max_len + 2, 11), 30)
        if "date" in header.lower() or header.lower() == "period":
            width = max(width, 13)
        ws.column_dimensions[get_column_letter(col_index)].width = width
    if ws.title == "Account ECL":
        ws.freeze_panes = "C2"
    if ws.title == "Controls":
        status_col = headers.index("status") + 1
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row, status_col)
            if cell.value == "PASS":
                cell.fill = PatternFill("solid", fgColor=PALE_GREEN)
                cell.font = Font(name="Aptos", size=10, bold=True, color=GREEN)


def _dashboard(wb: Workbook) -> None:
    stage_ws = wb["Stage Summary"]
    scen_ws = wb["Scenario Summary"]
    stage_rows = [[stage_ws.cell(r, c).value for c in range(1, 8)] for r in range(2, 5)]
    total_loans = sum(row[1] for row in stage_rows)
    total_exposure = sum(row[2] for row in stage_rows)
    total_ecl = sum(row[3] for row in stage_rows)
    ws = wb.create_sheet("Executive Dashboard", 0)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    ws["A1"] = "IFRS 9 Residential Mortgage Expected Credit Loss"
    ws["A1"].font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    for row in ws["A1:N2"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=NAVY)
    scenario_rows = {scen_ws.cell(row, 1).value: row for row in range(2, scen_ws.max_row + 1)}
    weight_text = " • ".join(f"{name} {scen_ws.cell(row, 2).value:.0%}" for name, row in scenario_rows.items())
    ws["A3"] = "Loan-level, scenario-weighted and discounted ECL | " + weight_text
    ws["A3"].font = Font(name="Aptos", size=10, italic=True, color=BLUE)
    ws["A3"].alignment = Alignment(horizontal="left")
    for cell in ws[3][:14]:
        cell.fill = PatternFill("solid", fgColor=PALE_BLUE)

    kpis = [("Key metric", "Value"), ("Mortgage accounts", total_loans),
            ("Gross exposure", total_exposure),
            ("Probability-weighted ECL", total_ecl),
            ("Portfolio coverage", total_ecl / total_exposure if total_exposure else 0)]
    for r, (label, value) in enumerate(kpis, start=5):
        ws.cell(r, 1, label)
        ws.cell(r, 2, value)
    for cell in ws[5]:
        if cell.column <= 2:
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.font = Font(bold=True, color=WHITE)
    for r in range(6, 10):
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=PALE_BLUE)
        ws.cell(r, 1).font = Font(bold=True, color=NAVY)
    ws["B7"].number_format = ws["B8"].number_format = '"$"#,##0'
    ws["B9"].number_format = "0.00%"

    stage_headers = ["Stage", "Loans", "Gross exposure", "ECL", "Coverage"]
    for c, value in enumerate(stage_headers, 1):
        ws.cell(12, c, value)
    for source_row, target_row in enumerate(range(13, 16), start=2):
        for col in range(1, 6):
            source_cell = stage_ws.cell(source_row, col).coordinate
            ws.cell(target_row, col, f"='Stage Summary'!{source_cell}")
    ws.cell(16, 1, "Total")
    ws["B16"] = total_loans
    ws["C16"] = total_exposure
    ws["D16"] = total_ecl
    ws["E16"] = total_ecl / total_exposure if total_exposure else 0

    scenario_headers = ["Scenario", "Weight", "Scenario ECL", "Weighted contribution"]
    for c, value in enumerate(scenario_headers, 7):
        ws.cell(12, c, value)
    scenario_map = [(name, source_row) for name, source_row in scenario_rows.items()]
    for row, (name, source_row) in enumerate(scenario_map, start=13):
        ws.cell(row, 7, name)
        for target_col, source_col in zip(range(8, 11), range(2, 5)):
            source_cell = scen_ws.cell(source_row, source_col).coordinate
            ws.cell(row, target_col, f"='Scenario Summary'!{source_cell}")
    ws["G16"] = "Total"
    ws["H16"] = "=SUM(H13:H15)"
    ws["I16"] = "Not additive"
    ws["J16"] = "=SUM(J13:J15)"

    for cell in list(ws[12])[:5] + list(ws[12])[6:10]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(bold=True, color=WHITE)
    thin = Side(style="thin", color="8EA9C1")
    for row in (16,):
        for col in list(range(1, 6)) + list(range(7, 11)):
            ws.cell(row, col).fill = PatternFill("solid", fgColor=PALE_GREEN)
            ws.cell(row, col).font = Font(bold=True, color=NAVY)
            ws.cell(row, col).border = Border(bottom=Side(style="double", color=NAVY))
    for row in range(13, 17):
        for col in (3, 4, 9, 10):
            ws.cell(row, col).number_format = '"$"#,##0'
        ws.cell(row, 5).number_format = "0.0000%"
        ws.cell(row, 8).number_format = "0%"

    note_lines = [
        "Method: Stage 1 sums discounted MPD × LGD × EAD for defaults in months 1–12; Stage 2 uses the remaining contractual life.",
        "Stage 3 is exposure less discounted scenario-specific expected recoveries and does not reuse the performing-loan PD curve.",
        "Each scenario ECL is calculated independently before applying scenario probabilities.",
    ]
    for row_number, note in enumerate(note_lines, start=34):
        ws.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=14)
        cell = ws.cell(row_number, 1, note)
        cell.font = Font(name="Aptos", size=10, color="7F6000")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws["A34:N37"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=PALE_AMBER)

    chart = BarChart()
    chart.type = "col"
    chart.title = "Gross Exposure by Stage (USD)"
    chart.y_axis.numFmt = '$#,##0'
    chart.add_data(Reference(stage_ws, min_col=3, max_col=3, min_row=1, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(stage_ws, min_col=1, min_row=2, max_row=4))
    chart.legend = None
    chart.height, chart.width = 7, 11
    ws.add_chart(chart, "A18")

    chart2 = BarChart()
    chart2.type = "col"
    chart2.title = "Scenario ECL Sensitivity (USD)"
    chart2.y_axis.numFmt = '$#,##0'
    chart2.add_data(Reference(scen_ws, min_col=3, max_col=3, min_row=1, max_row=4), titles_from_data=True)
    chart2.set_categories(Reference(scen_ws, min_col=1, min_row=2, max_row=4))
    chart2.legend = None
    chart2.height, chart2.width = 7, 8
    ws.add_chart(chart2, "H18")

    for column, width in {"A": 30, "B": 18, "C": 16, "D": 16, "E": 16, "G": 16, "H": 18, "I": 18, "J": 20}.items():
        ws.column_dimensions[column].width = width
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 28
    for row_number in range(34, 37):
        ws.row_dimensions[row_number].height = 20


def _sources_sheet(wb: Workbook, cfg: dict) -> None:
    ws = wb.create_sheet("Sources & Assumptions")
    rows = [
        ("Item", "Value / source"),
        ("Project", cfg["project"]["name"]),
        ("Freddie Mac source", "https://www.freddiemac.com/research/datasets/sf-loanlevel-dataset"),
        ("FRED unemployment", "https://fred.stlouisfed.org/series/UNRATE"),
        ("FRED real GDP", "https://fred.stlouisfed.org/series/GDPC1"),
        ("FHFA HPI", "https://www.fhfa.gov/data/hpi"),
        ("Reporting date", cfg["reporting"]["reporting_date"]),
        ("Scenario weights", "Selected by the configured blend of management assumptions and historical macro-regime frequencies"),
        ("Default", "90+ DPD, REO, or configured credit-related terminal event"),
        ("EIR approximation", cfg["discounting"]["eir_approximation"]),
        ("Synthetic source data", "None; model estimates and illustrative scenarios are separately flagged"),
    ]
    _write_frame(ws, pd.DataFrame(rows[1:], columns=rows[0]))
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 92
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")


def build_excel_report(cfg: dict) -> Path:
    table_dir = Path(cfg["paths"]["tables"])
    output = Path(cfg["paths"]["report"])
    wb = Workbook()
    wb.remove(wb.active)
    for file_name, sheet_name in SHEETS:
        frame = pd.read_csv(table_dir / file_name)
        for col in frame.columns:
            if col == "period" or col.endswith("_date"):
                converted = pd.to_datetime(frame[col], errors="coerce")
                if converted.notna().any():
                    frame[col] = converted
        ws = wb.create_sheet(sheet_name)
        _write_frame(ws, frame)
    _dashboard(wb)
    _sources_sheet(wb, cfg)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output
